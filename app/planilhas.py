"""Leitura e processamento das planilhas.

- Tabela do gerente: código Sankhya + preços CIF e FOB.
- Exportação do Mercos: catálogo de uma filial, no modelo de importação do Mercos.

Fluxo: ler_tabela_gerente -> ler_exportacao_mercos -> montar_plano (limpeza)
       -> aplicar_precos -> gerar_importacao_xlsx / gerar_relatorio_xlsx
"""
import difflib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class PlanilhaInvalida(Exception):
    """Erro com mensagem pronta para mostrar ao usuário."""


# ---------------------------------------------------------------- utilidades

def normalizar_texto(valor) -> str:
    """Maiúsculas, sem acento e sem espaços repetidos (para comparar nomes e cabeçalhos)."""
    if valor is None:
        return ""
    texto = unicodedata.normalize("NFKD", str(valor)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().upper()


def normalizar_codigo(valor) -> str:
    """Código como texto: 10001, 10001.0 e ' 10001 ' viram '10001'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if re.fullmatch(r"\d+\.0+", texto):
        texto = texto.split(".")[0]
    return texto.upper()


def numero(valor) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def arred(valor: float | None) -> float | None:
    """Arredondamento comercial em 2 casas (meio centavo sobe: 277,155 -> 277,16)."""
    if valor is None:
        return None
    return float(Decimal(repr(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def primeira_linha_cabecalho(valor) -> str:
    """Cabeçalhos do Mercos vêm como 'Código do produto\\n(recomendado)'; usa só a 1ª linha."""
    return normalizar_texto(str(valor).split("\n")[0]) if valor is not None else ""


def ncm_completo(valor) -> bool:
    return bool(re.fullmatch(r"\d{8}", re.sub(r"\D", "", str(valor or ""))))


def semelhanca(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, normalizar_texto(a), normalizar_texto(b)).ratio()


def abrir_pasta(dados: bytes):
    """Abre o .xlsx duas vezes: com fórmulas (para análise) e com valores calculados."""
    try:
        com_formulas = openpyxl.load_workbook(io.BytesIO(dados), data_only=False)
        com_valores = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
    except Exception as erro:  # arquivo corrompido, .xls antigo, outro formato
        raise PlanilhaInvalida(
            "Não consegui abrir o arquivo como planilha .xlsx. "
            "Se for .xls (Excel antigo) ou .csv, abra no Excel e salve como .xlsx."
        ) from erro
    return com_formulas, com_valores


def analisar_pasta(pasta_formulas) -> dict:
    """Raio-x da planilha para o log: abas, linhas/colunas ocultas, filtros, vínculos externos."""
    abas = []
    for aba in pasta_formulas.worksheets:
        linhas_ocultas = sum(1 for d in aba.row_dimensions.values() if d.hidden)
        colunas_ocultas = sorted(
            {get_column_letter(c) for d in aba.column_dimensions.values() if d.hidden
             for c in range(d.min or 1, (d.max or d.min or 1) + 1)}
        )
        formulas = 0
        refs_externas = 0
        for linha in aba.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    formulas += 1
                    if "[" in celula.value:
                        refs_externas += 1
        abas.append({
            "nome": aba.title,
            "estado": aba.sheet_state,
            "dimensao": aba.dimensions,
            "linhas": aba.max_row,
            "colunas": aba.max_column,
            "linhas_ocultas": linhas_ocultas,
            "colunas_ocultas": colunas_ocultas,
            "filtro": aba.auto_filter.ref,
            "formulas": formulas,
            "formulas_com_vinculo_externo": refs_externas,
        })
    vinculos = []
    for vinculo in getattr(pasta_formulas, "_external_links", []) or []:
        try:
            vinculos.append(vinculo.file_link.Target)
        except AttributeError:
            vinculos.append("?")
    return {"abas": abas, "vinculos_externos": vinculos}


# ---------------------------------------------------------- tabela do gerente

@dataclass
class ItemTabela:
    codigo: str
    nome: str
    embalagem: str
    cif: float | None
    fob: float | None
    linha: int


@dataclass
class TabelaLida:
    itens: dict[str, ItemTabela]
    conflitos: list[dict]
    repetidos_iguais: list[str]
    sem_preco: list[dict]
    aba: str
    linha_cabecalho: int
    colunas: dict
    analise: dict


def _achar_cabecalho_tabela(aba):
    for numero_linha in range(1, min(aba.max_row, 40) + 1):
        celulas = {c.column: normalizar_texto(c.value) for c in aba[numero_linha] if c.value is not None}
        col_codigo = next((col for col, t in celulas.items() if t.startswith("CODIGO")), None)
        col_cif = next((col for col, t in celulas.items() if "TABELA" in t and "CIF" in t), None) or \
            next((col for col, t in celulas.items() if "CIF" in t), None)
        col_fob = next((col for col, t in celulas.items() if "TABELA" in t and "FOB" in t), None) or \
            next((col for col, t in celulas.items() if "FOB" in t), None)
        if col_codigo and col_cif and col_fob:
            col_nome = next((col for col, t in celulas.items() if "PRODUTO" in t and col != col_codigo), None) or \
                next((col for col, t in celulas.items() if "NOME" in t or "DESCR" in t), None)
            col_emb = next((col for col, t in celulas.items() if "EMBALAGEM" in t), None)
            return numero_linha, {"codigo": col_codigo, "nome": col_nome, "embalagem": col_emb,
                                  "cif": col_cif, "fob": col_fob}
    return None, None


def ler_tabela_gerente(dados: bytes) -> TabelaLida:
    pasta_f, pasta_v = abrir_pasta(dados)
    analise = analisar_pasta(pasta_f)

    for aba_v in pasta_v.worksheets:
        linha_cab, colunas = _achar_cabecalho_tabela(aba_v)
        if linha_cab:
            break
    else:
        raise PlanilhaInvalida(
            "Não encontrei na tabela do gerente uma linha de cabeçalho com 'Código', "
            "'... CIF' e '... FOB'. Confira se é a planilha certa."
        )

    aba_f = pasta_f[aba_v.title]
    itens: dict[str, ItemTabela] = {}
    conflitos, repetidos_iguais, sem_preco = [], [], []
    formulas_sem_valor = 0

    for numero_linha in range(linha_cab + 1, aba_v.max_row + 1):
        codigo = normalizar_codigo(aba_v.cell(numero_linha, colunas["codigo"]).value)
        if not re.fullmatch(r"\d+", codigo):
            continue  # títulos de seção ("Graxas Litio"), linhas vazias, observações
        cif = numero(aba_v.cell(numero_linha, colunas["cif"]).value)
        fob = numero(aba_v.cell(numero_linha, colunas["fob"]).value)
        for col in (colunas["cif"], colunas["fob"]):
            bruto = aba_f.cell(numero_linha, col).value
            if isinstance(bruto, str) and bruto.startswith("=") and aba_v.cell(numero_linha, col).value is None:
                formulas_sem_valor += 1
        nome = str(aba_v.cell(numero_linha, colunas["nome"]).value or "").strip() if colunas["nome"] else ""
        embalagem = str(aba_v.cell(numero_linha, colunas["embalagem"]).value or "").strip() if colunas["embalagem"] else ""
        item = ItemTabela(codigo, nome, embalagem, arred(cif), arred(fob), numero_linha)

        if not cif or not fob or cif <= 0 or fob <= 0:
            sem_preco.append({"codigo": codigo, "nome": nome, "linha": numero_linha, "cif": cif, "fob": fob})

        anterior = itens.get(codigo)
        if anterior is None:
            itens[codigo] = item
        elif (anterior.cif, anterior.fob) == (item.cif, item.fob):
            repetidos_iguais.append(codigo)
        else:
            conflitos.append({
                "codigo": codigo,
                "linhas": [anterior.linha, numero_linha],
                "nomes": [anterior.nome, nome],
                "precos": [(anterior.cif, anterior.fob), (item.cif, item.fob)],
            })

    if formulas_sem_valor:
        raise PlanilhaInvalida(
            f"{formulas_sem_valor} preços são fórmulas sem valor calculado. "
            "Abra a planilha no Excel, deixe calcular, salve e envie de novo."
        )
    if not itens:
        raise PlanilhaInvalida("A tabela do gerente não tem nenhum produto com código numérico.")

    # código com conflito não pode ser usado até o gerente corrigir
    for conflito in conflitos:
        itens.pop(conflito["codigo"], None)

    return TabelaLida(itens, conflitos, sorted(set(repetidos_iguais)), sem_preco, aba_v.title, linha_cab,
                      {k: get_column_letter(v) if v else None for k, v in colunas.items()}, analise)


# ------------------------------------------------------- exportação do Mercos

CABECALHOS_MERCOS = {
    "codigo": lambda t: t == "CODIGO DO PRODUTO",
    "nome": lambda t: t == "NOME DO PRODUTO",
    "preco_tabela": lambda t: t == "PRECO DE TABELA",
    "preco_minimo": lambda t: t == "PRECO MINIMO",
    "ncm": lambda t: t == "NCM",
    "estoque": lambda t: t.startswith("QUANTIDADE EM ESTOQUE"),
    "ativo": lambda t: t.startswith("ATIVO / INATIVO"),
    "exibido": lambda t: t.startswith("EXIBIDO"),
    "cif": lambda t: "CIF" in t,
    "fob": lambda t: "FOB" in t,
}
OBRIGATORIOS = ("codigo", "nome", "preco_tabela", "preco_minimo", "cif", "fob")


@dataclass
class LinhaMercos:
    linha: int
    codigo: str
    nome: str
    valores: list
    oculta: bool

    def valor(self, colunas: dict, campo: str):
        indice = colunas.get(campo)
        return self.valores[indice - 1] if indice else None


@dataclass
class ExportacaoLida:
    cabecalho: list
    colunas: dict
    linhas: list[LinhaMercos]
    analise: dict
    total_colunas: int


def ler_exportacao_mercos(dados: bytes) -> ExportacaoLida:
    pasta_f, pasta_v = abrir_pasta(dados)
    analise = analisar_pasta(pasta_f)
    aba = pasta_v.worksheets[0]
    total_colunas = aba.max_column
    cabecalho = [aba.cell(1, c).value for c in range(1, total_colunas + 1)]

    colunas = {}
    for indice, valor in enumerate(cabecalho, start=1):
        texto = primeira_linha_cabecalho(valor)
        for campo, combina in CABECALHOS_MERCOS.items():
            if campo not in colunas and texto and combina(texto):
                colunas[campo] = indice
                break
    faltando = [c for c in OBRIGATORIOS if c not in colunas]
    if faltando:
        nomes = {"codigo": "Código do produto", "nome": "Nome do produto", "preco_tabela": "Preço de Tabela",
                 "preco_minimo": "Preço Mínimo", "cif": "tabela de preços ... CIF", "fob": "tabela de preços ... FOB"}
        raise PlanilhaInvalida(
            "Este arquivo não parece uma exportação de produtos do Mercos. Faltam as colunas: "
            + ", ".join(nomes[c] for c in faltando) + "."
        )

    ocultas = {n for n, d in pasta_f.worksheets[0].row_dimensions.items() if d.hidden}
    linhas = []
    for numero_linha, valores in enumerate(aba.iter_rows(min_row=2, max_col=total_colunas, values_only=True), start=2):
        valores = list(valores)
        codigo = normalizar_codigo(valores[colunas["codigo"] - 1])
        nome = str(valores[colunas["nome"] - 1] or "").strip()
        if not codigo and not nome:
            continue
        linhas.append(LinhaMercos(numero_linha, codigo, nome, valores, numero_linha in ocultas))
    if not linhas:
        raise PlanilhaInvalida("A exportação do Mercos não tem nenhum produto.")
    return ExportacaoLida(cabecalho, colunas, linhas, analise, total_colunas)


def foto_catalogo(exportacao: ExportacaoLida) -> dict:
    codigos = [l.codigo for l in exportacao.linhas if l.codigo]
    contagem = {}
    for c in codigos:
        contagem[c] = contagem.get(c, 0) + 1
    inativos = sum(1 for l in exportacao.linhas if str(l.valor(exportacao.colunas, "ativo") or "").strip() == "1")
    return {
        "total_linhas": len(exportacao.linhas),
        "com_codigo": len(codigos),
        "sem_codigo": sum(1 for l in exportacao.linhas if not l.codigo),
        "duplicados": sum(n - 1 for n in contagem.values() if n > 1),
        "inativos": inativos,
        "ocultas": sum(1 for l in exportacao.linhas if l.oculta),
        "codigos": sorted(set(codigos)),
    }


# ------------------------------------------------------------ plano de limpeza

@dataclass
class Decisao:
    linha: int
    codigo: str
    nome: str
    acao: str               # manter | remover | preencher
    motivo: str
    grupo: str              # unico | duplicado | sem_codigo
    manual: bool = False
    codigo_sugerido: str | None = None


@dataclass
class Plano:
    decisoes: list[Decisao]
    avisos: list[dict] = field(default_factory=list)


def _completude(linha: LinhaMercos, colunas: dict) -> float:
    pontos = 0.0
    if (numero(linha.valor(colunas, "preco_tabela")) or 0) > 0:
        pontos += 2
    if ncm_completo(linha.valor(colunas, "ncm")):
        pontos += 1
    if "ENVASE" not in normalizar_texto(linha.nome):
        pontos += 1
    if linha.valor(colunas, "estoque") not in (None, ""):
        pontos += 0.5
    return pontos


def montar_plano(exportacao: ExportacaoLida, tabela: dict[str, ItemTabela]) -> Plano:
    """Regras de limpeza, ajustadas ao comportamento real da importação do Mercos ("Atualizar"):
    - o Mercos casa os produtos pelo código e CRIA os que não existem;
    - se o código aparece mais de uma vez no catálogo, ele APAGA todas as cópias e cria um
      cadastro novo — que nasce sem vínculo com o connector (teste de 06/10/2026).
    Por isso:
    1. compara por código (texto, sem diferenciar maiúsculas/espaços);
    2. código único fica (mesmo fora da tabela do gerente);
    3. código repetido NÃO vai no arquivo: é listado para inativar a cópia errada no Mercos;
    4. linha sem código nunca vai no arquivo (mandar com código criaria produto novo);
    5. avisos de nomes trocados e de nomes que só diferem em maiúsculas/espaços.
    """
    colunas = exportacao.colunas
    decisoes: list[Decisao] = []
    avisos: list[dict] = []

    por_codigo: dict[str, list[LinhaMercos]] = {}
    for linha in exportacao.linhas:
        if linha.codigo:
            por_codigo.setdefault(linha.codigo, []).append(linha)

    for codigo, grupo in por_codigo.items():
        item = tabela.get(codigo)
        if len(grupo) == 1:
            linha = grupo[0]
            decisoes.append(Decisao(linha.linha, codigo, linha.nome, "manter", "código único", "unico"))
            continue

        # aponta qual cópia parece a certa, só como orientação para quem for inativar no Mercos
        referencia = f"{item.nome} {item.embalagem}" if item else None
        pontuadas = sorted(
            ((semelhanca(l.nome, referencia) if referencia else 0.0, _completude(l, colunas), -l.linha, l) for l in grupo),
            key=lambda p: (p[0], p[1], p[2]), reverse=True)
        melhor = pontuadas[0][3]
        for linha in grupo:
            decisoes.append(Decisao(
                linha.linha, codigo, linha.nome, "remover",
                f"código repetido no Mercos ({len(grupo)}x) — não vai no arquivo",
                "duplicado"))
        avisos.append({
            "tipo": "duplicado_mercos",
            "codigo": codigo,
            "nome": melhor.nome,
            "mensagem": f"O código {codigo} tem {len(grupo)} cadastros no Mercos (linhas "
                        f"{', '.join(str(l.linha) for l in grupo)} da exportação) e ficou FORA do arquivo: "
                        "importar apagaria as cópias e criaria um cadastro novo sem vínculo com o connector. "
                        f"Inative no Mercos a cópia errada (a mais completa parece ser a da linha {melhor.linha}) "
                        "e gere a importação de novo.",
        })
        for linha in grupo:
            if linha is not melhor and semelhanca(linha.nome, melhor.nome) < 0.5:
                avisos.append({
                    "tipo": "nome_trocado",
                    "codigo": codigo,
                    "nome": linha.nome,
                    "mensagem": f"O código {codigo} também aparece com o nome '{linha.nome}' "
                                f"(o certo parece ser '{melhor.nome}'). Cadastro errado no Mercos.",
                })

    nomes_tabela: dict[str, list[str]] = {}
    for item in tabela.values():
        for nome in {normalizar_texto(item.nome), normalizar_texto(f"{item.nome} {item.embalagem}")}:
            if nome:
                nomes_tabela.setdefault(nome, []).append(item.codigo)
    for linha in exportacao.linhas:
        if linha.codigo:
            continue
        candidatos = sorted(set(nomes_tabela.get(normalizar_texto(linha.nome), [])))
        motivo = "sem código no Mercos — não vai no arquivo"
        if len(candidatos) == 1:
            motivo += f" (o nome corresponde ao código {candidatos[0]}; inative este cadastro no Mercos)"
        decisoes.append(Decisao(linha.linha, "", linha.nome, "remover", motivo, "sem_codigo"))

    for codigo, grupo in por_codigo.items():
        item = tabela.get(codigo)
        if not item:
            continue
        for linha in grupo:
            if linha.nome.strip() != item.nome.strip() and normalizar_texto(linha.nome) == normalizar_texto(item.nome):
                avisos.append({"tipo": "nome_formatacao", "codigo": codigo, "nome": linha.nome,
                               "mensagem": f"Nome difere do gerente só em maiúsculas/espaços: '{linha.nome}' x '{item.nome}'."})

    decisoes.sort(key=lambda d: d.linha)
    return Plano(decisoes, avisos)


# ----------------------------------------------------------- aplicar preços

@dataclass
class LinhaSaida:
    codigo: str
    nome: str
    valores: list
    cif_antes: float | None
    fob_antes: float | None
    cif_depois: float | None
    fob_depois: float | None
    preco_tabela: float | None
    preco_minimo: float | None
    origem: str               # tabela | mantido
    inativo: bool
    alterado: bool = True     # False = nada muda no Mercos; a linha não vai no arquivo

    @property
    def var_cif(self) -> float | None:
        if not self.cif_antes or self.cif_depois is None:
            return None
        return round((self.cif_depois - self.cif_antes) / self.cif_antes * 100, 2)


def aplicar_precos(exportacao: ExportacaoLida, decisoes: list[Decisao], tabela: dict[str, ItemTabela],
                   regra_preco_minimo: str = "fob", desconto_max: float | None = None) -> list[LinhaSaida]:
    colunas = exportacao.colunas
    por_linha = {l.linha: l for l in exportacao.linhas}
    saida = []
    for decisao in decisoes:
        if decisao.acao == "remover":
            continue
        original = por_linha[decisao.linha]
        valores = list(original.valores)
        codigo = decisao.codigo_sugerido if decisao.acao == "preencher" else original.codigo
        if decisao.acao == "preencher":
            valores[colunas["codigo"] - 1] = codigo

        cif_antes = arred(numero(original.valor(colunas, "cif")))
        fob_antes = arred(numero(original.valor(colunas, "fob")))
        item = tabela.get(codigo)
        if item and item.cif and item.fob:
            cif, fob, origem = item.cif, item.fob, "tabela"
        else:
            # fora da tabela do gerente: tudo exatamente como está no Mercos (inclusive campos vazios);
            # o desconto do gerente não vale aqui
            cif, fob, origem = cif_antes, fob_antes, "mantido"

        precos = [p for p in (cif, fob) if p is not None]
        preco_tabela = max(precos) if precos else None
        if origem == "mantido":
            preco_tabela = arred(numero(original.valor(colunas, "preco_tabela")))
            preco_minimo = arred(numero(original.valor(colunas, "preco_minimo")))
        elif regra_preco_minimo == "desconto" and desconto_max is not None and preco_tabela:
            preco_minimo = arred(preco_tabela * (1 - desconto_max / 100))
        else:
            preco_minimo = min(precos) if precos else None
        def ativo_original(campo):
            valor = str(original.valor(colunas, campo) or "0").strip()
            return 1 if valor in ("1", "1.0") else 0

        if origem == "mantido":
            # fora da tabela do gerente: tudo fica como está, inclusive ativo/inativo
            inativo = "ativo" in colunas and ativo_original("ativo") == 1
        else:
            inativo = not preco_tabela or not preco_minimo or preco_tabela <= 0 or preco_minimo <= 0

        valores[colunas["cif"] - 1] = cif
        valores[colunas["fob"] - 1] = fob
        valores[colunas["preco_tabela"] - 1] = preco_tabela
        valores[colunas["preco_minimo"] - 1] = preco_minimo
        # Estoque vai em branco: o Mercos mantém o atual (confirmado pelo suporte e em teste) e quem
        # atualiza é o connector. NCM NÃO pode ir em branco: o Mercos apaga o campo (teste de 06/10/2026),
        # então ele segue exatamente como veio da exportação.
        if "estoque" in colunas:
            valores[colunas["estoque"] - 1] = None
        for campo in ("ativo", "exibido"):
            if campo in colunas:
                valores[colunas[campo] - 1] = (ativo_original(campo) if origem == "mantido"
                                               else (1 if inativo else 0))

        antes = (arred(numero(original.valor(colunas, "preco_tabela"))), arred(numero(original.valor(colunas, "preco_minimo"))),
                 cif_antes, fob_antes, *(ativo_original(c) for c in ("ativo", "exibido") if c in colunas))
        depois = (preco_tabela, preco_minimo, cif, fob,
                  *(valores[colunas[c] - 1] for c in ("ativo", "exibido") if c in colunas))

        saida.append(LinhaSaida(codigo, original.nome, valores, cif_antes, fob_antes, cif, fob,
                                preco_tabela, preco_minimo, origem, inativo, alterado=antes != depois))
    return saida


def somente_alterados(saida: list[LinhaSaida]) -> list[LinhaSaida]:
    """O arquivo leva só o que muda: cada linha enviada é um risco de efeito colateral no Mercos
    (ex.: código com cópia inativa escondida, que o Mercos apagaria e recriaria)."""
    return [s for s in saida if s.alterado]


# ------------------------------------------------------------ gerar arquivos

def _estilizar_cabecalho(aba, total_colunas: int) -> None:
    for c in range(1, total_colunas + 1):
        celula = aba.cell(1, c)
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        aba.column_dimensions[get_column_letter(c)].width = 18
    aba.freeze_panes = "A2"


def gerar_importacao_xlsx(exportacao: ExportacaoLida, saida: list[LinhaSaida]) -> bytes:
    """Arquivo no mesmo modelo da exportação do Mercos, só com valores (sem fórmulas, filtros ou ocultas)."""
    pasta = openpyxl.Workbook()
    aba = pasta.active
    aba.title = "Planilha1"
    aba.append(exportacao.cabecalho)
    for linha in saida:
        aba.append(linha.valores)
    _estilizar_cabecalho(aba, exportacao.total_colunas)
    aba.column_dimensions["B"].width = 50
    arquivo = io.BytesIO()
    pasta.save(arquivo)
    return arquivo.getvalue()


def _aba_tabela(pasta, titulo: str, cabecalho: list, linhas: list, destaque=None) -> None:
    aba = pasta.create_sheet(titulo[:31])
    aba.append(cabecalho)
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for linha in linhas:
        aba.append(linha)
        if destaque and destaque(linha):
            for celula in aba[aba.max_row]:
                celula.fill = amarelo
    for c in range(1, len(cabecalho) + 1):
        aba.cell(1, c).font = Font(bold=True)
        aba.column_dimensions[get_column_letter(c)].width = 16
    if len(cabecalho) > 1:
        aba.column_dimensions["B"].width = 55
    aba.freeze_panes = "A2"
    if linhas:
        aba.auto_filter.ref = aba.dimensions


def gerar_relatorio_xlsx(resumo: dict, plano: Plano, saida: list[LinhaSaida], cadastrar: list[dict],
                         variacao_destaque: float) -> bytes:
    pasta = openpyxl.Workbook()
    aba = pasta.active
    aba.title = "Resumo"
    for chave, valor in resumo.items():
        aba.append([chave, valor])
    aba.column_dimensions["A"].width = 45
    aba.column_dimensions["B"].width = 60

    alterados = [s for s in saida if s.origem == "tabela" and s.alterado]
    _aba_tabela(
        pasta, "Preços alterados",
        ["Código", "Produto", "CIF antes", "CIF depois", "Variação CIF %", "FOB antes", "FOB depois", "Preço mínimo"],
        [[s.codigo, s.nome, s.cif_antes, s.cif_depois, s.var_cif, s.fob_antes, s.fob_depois, s.preco_minimo]
         for s in alterados],
        destaque=lambda l: l[4] is not None and abs(l[4]) >= variacao_destaque,
    )
    _aba_tabela(
        pasta, "Fora da tabela (mantidos)",
        ["Código", "Produto", "Preço de tabela mantido", "Preço mínimo mantido", "Situação", "Vai no arquivo"],
        [[s.codigo, s.nome, s.preco_tabela, s.preco_minimo, "INATIVO" if s.inativo else "ativo",
          "sim" if s.alterado else "não (nada muda)"] for s in saida if s.origem == "mantido"],
    )
    _aba_tabela(
        pasta, "Sem alteração (não enviados)",
        ["Código", "Produto", "Preço de tabela", "Preço mínimo"],
        [[s.codigo, s.nome, s.preco_tabela, s.preco_minimo] for s in saida if not s.alterado],
    )
    _aba_tabela(
        pasta, "Limpeza",
        ["Linha", "Código", "Produto", "Ação", "Motivo", "Decisão manual"],
        [[d.linha, d.codigo or d.codigo_sugerido or "", d.nome, d.acao, d.motivo, "sim" if d.manual else ""]
         for d in plano.decisoes if d.acao != "manter" or d.grupo != "unico"],
    )
    _aba_tabela(
        pasta, "Cadastrar no Sankhya",
        ["Código", "Produto (tabela do gerente)", "Embalagem", "CIF", "FOB"],
        [[c["codigo"], c["nome"], c["embalagem"], c["cif"], c["fob"]] for c in cadastrar],
    )
    _aba_tabela(pasta, "Avisos", ["Tipo", "Código", "Mensagem"],
                [[a["tipo"], a["codigo"], a["mensagem"]] for a in plano.avisos])
    arquivo = io.BytesIO()
    pasta.save(arquivo)
    return arquivo.getvalue()
